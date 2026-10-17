"""Word (.docx) export, built from the same report context as the HTML/PDF templates.

Needs python-docx (pip install python-docx). Word can't use the web fonts embedded in the HTML report,
so each template maps to fonts that ship with Microsoft Office / Windows.
"""
from __future__ import annotations

import base64
import copy
import io
import re
from html.parser import HTMLParser

from .config import ReportConfig
from .parser import ScanReport
from .render import build_context, resolve_template
from .sanitize import sanitize_html

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.oxml.table import CT_Tbl
    from docx.shared import Cm, Pt, RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from lxml import etree
    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover - reported to the user instead
    DOCX_AVAILABLE = False

DOCX_MISSING = "Word export needs the python-docx package: pip install python-docx"

THEMES = {
    "corporate": {"accent": "1F3A5F", "heading": "Cambria", "body": "Calibri", "numbered": True},
    "modern": {"accent": "5B3CC4", "heading": "Segoe UI Semibold", "body": "Segoe UI", "numbered": False},
    "minimal": {"accent": "111111", "heading": "Arial", "body": "Arial", "numbered": False},
}
# Office-safe stand-ins for the web fonts offered in the font picker.
FONT_EQUIVALENTS = {
    "inter": "Segoe UI", "source sans 3": "Calibri", "source serif 4": "Cambria",
    "ibm plex sans": "Arial", "jetbrains mono": "Consolas", "system-ui": "Segoe UI",
}
GENERIC_FAMILIES = {"sans-serif", "serif", "monospace", "-apple-system", "cursive", "fantasy"}
MONO = "Consolas"
TEXT, MUTED, SURFACE, BORDER = "1D2330", "5F6B7A", "F4F6F9", "DDE3EA"
CODE_BG = "F3F5F8"
SEV_COLORS = {"High": "C62828", "Medium": "E46C0A", "Low": "D6A400", "Information": "4F7FB5",
              "False positive": "8A8F98"}
SEV_TEXT = {"Low": "2B2100"}
SEV_SHORT = {"Information": "Info", "False positive": "FP"}   # narrow table headers
_XML_INVALID = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


def _clean(text) -> str:
    return _XML_INVALID.sub("", str(text or ""))


def _hex(color: str) -> str | None:
    m = re.fullmatch(r"#?([0-9a-fA-F]{6}|[0-9a-fA-F]{3})", (color or "").strip())
    if not m:
        return None
    h = m.group(1)
    return (h if len(h) == 6 else "".join(c * 2 for c in h)).upper()


def _font_from_stack(stack: str) -> str | None:
    for name in (stack or "").split(","):
        name = name.strip().strip("'\"")
        if not name or name.lower() in GENERIC_FAMILIES:
            continue
        return FONT_EQUIVALENTS.get(name.lower(), name)
    return None


def _theme(cfg: ReportConfig) -> dict:
    try:
        key = resolve_template(cfg.template).name
    except FileNotFoundError:
        key = "corporate"
    theme = dict(THEMES.get(key, THEMES["corporate"]))
    theme["accent"] = _hex(cfg.accent_color) or theme["accent"]
    custom = _font_from_stack(cfg.font_family)
    if custom:
        theme["body"] = custom
        if theme["heading"] != "Cambria" or key != "corporate":
            theme["heading"] = custom
    return theme


# ---------------------------------------------------------------- low-level helpers

def _rgb(h: str) -> "RGBColor":
    return RGBColor.from_string(h)


def _apply_font(rpr, name: str) -> None:
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    # Theme font attributes (used by Word's built-in heading styles) win over named fonts, so drop them.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)


def _set_font(style, name: str) -> None:
    style.font.name = name
    _apply_font(style.element.get_or_add_rPr(), name)


def _run_font(run, name: str) -> None:
    run.font.name = name
    _apply_font(run._element.get_or_add_rPr(), name)


def _widths(table, widths_cm: list[float]) -> None:
    """Fixed column widths (Word ignores cell widths while the table autofits)."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        col.set(qn("w:w"), str(int(Cm(w).twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def _spread(fixed: list[float], flexible: int, total: float = 16.6) -> list[float]:
    """Widths for a table whose first columns are fixed and the remaining `flexible` share the rest."""
    rest = max(total - sum(fixed), 1.2 * flexible)
    return fixed + [rest / flexible] * flexible


def _shade(cell, fill: str) -> None:
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _run_shade(run, fill: str) -> None:
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    rpr.append(shd)


def _cell_margins(table, top=60, bottom=60, left=100, right=100) -> None:
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _borders(table, color: str = BORDER, inside: bool = True, size: int = 4) -> None:
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ()):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if color else "nil")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color or "auto")
        borders.append(el)
    tblpr.append(borders)


def _no_borders(table) -> None:
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _no_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    """Insert a Word field (PAGE, NUMPAGES, TOC ...)."""
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el
    r1 = paragraph.add_run()
    r1._element.append(fld("begin"))
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._element.append(it)
    r3 = paragraph.add_run()
    r3._element.append(fld("separate"))
    paragraph.add_run(placeholder)
    r4 = paragraph.add_run()
    r4._element.append(fld("end"))


def _bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _hyperlink(paragraph, text: str, url: str = "", anchor: str = "", color: str = "0563C1", font: str | None = None,
               bold: bool = False):
    link = OxmlElement("w:hyperlink")
    if anchor:
        link.set(qn("w:anchor"), anchor)
    else:
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if font:
        fonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), font)
        rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = _clean(text)
    run.append(t)
    link.append(run)
    paragraph._p.append(link)
    return link


def _keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- Burp HTML → paragraphs

class _HtmlToDocx(HTMLParser):
    """Writes sanitized Burp HTML (p, lists, bold/italic, code, links, tables...) into a document or cell."""

    BLOCKS = {"p", "div", "blockquote", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "tr"}

    def __init__(self, container, body_font: str):
        super().__init__(convert_charrefs=True)
        self.c = container
        self.body_font = body_font
        self.para = None
        self.fmt: list[str] = []
        self.lists: list[list] = []      # [kind, counter]
        self.link: str | None = None
        self.link_text: list[str] = []
        self.pre = 0
        self.first_cell = True

    def _new_para(self, style: str | None = None):
        self.para = self.c.add_paragraph(style=style) if style else self.c.add_paragraph()
        self.para.paragraph_format.space_after = Pt(4)
        return self.para

    def _ensure(self):
        return self.para or self._new_para()

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in self.BLOCKS:
            self._new_para()
            if tag.startswith("h") or tag == "th":
                self.fmt.append("b")
            if tag == "pre":
                self.pre += 1
                self.fmt.append("code")
            if tag == "blockquote":
                self.para.paragraph_format.left_indent = Cm(0.8)
            self.first_cell = True
        elif tag in ("ul", "ol"):
            self.lists.append([tag, 0])
            self.para = None
        elif tag == "li":
            kind = self.lists[-1] if self.lists else ["ul", 0]
            depth = len(self.lists)
            if kind[0] == "ol":
                # Explicit numbers: Word's List Number style would continue counting across separate lists.
                kind[1] += 1
                self._new_para()
                self.para.paragraph_format.left_indent = Cm(0.63 * depth)
                self.para.paragraph_format.first_line_indent = Cm(-0.5)
                self.para.add_run(f"{kind[1]}. ")
            else:
                self._new_para("List Bullet" if depth <= 1 else "List Bullet 2")
            self.para.paragraph_format.space_after = Pt(2)
        elif tag == "br":
            self._ensure().add_run().add_break()
        elif tag == "hr":
            self._new_para()
        elif tag in ("td", "th"):
            if not self.first_cell:
                self._ensure().add_run("  |  ")
            self.first_cell = False
            if tag == "th":
                self.fmt.append("b")
        elif tag in ("b", "strong"):
            self.fmt.append("b")
        elif tag in ("i", "em"):
            self.fmt.append("i")
        elif tag == "u":
            self.fmt.append("u")
        elif tag == "code":
            self.fmt.append("code")
        elif tag in ("sub", "sup"):
            self.fmt.append(tag)
        elif tag == "a" and a.get("href"):
            self.link, self.link_text = a["href"], []

    def handle_endtag(self, tag):
        if tag in self.BLOCKS:
            if tag.startswith("h") and "b" in self.fmt:
                self.fmt.remove("b")
            if tag == "pre":
                self.pre = max(0, self.pre - 1)
                if "code" in self.fmt:
                    self.fmt.remove("code")
            self.para = None
        elif tag in ("ul", "ol"):
            if self.lists:
                self.lists.pop()
            self.para = None
        elif tag == "li":
            self.para = None
        elif tag == "th" and "b" in self.fmt:
            self.fmt.remove("b")
        elif tag in ("b", "strong") and "b" in self.fmt:
            self.fmt.remove("b")
        elif tag in ("i", "em") and "i" in self.fmt:
            self.fmt.remove("i")
        elif tag in ("u", "code", "sub", "sup") and tag in self.fmt:
            self.fmt.remove(tag)
        elif tag == "a" and self.link:
            text = "".join(self.link_text).strip() or self.link
            _hyperlink(self._ensure(), text, url=self.link)
            self.link = None

    def handle_data(self, data):
        if not self.pre:
            data = re.sub(r"\s+", " ", data)
            if not data.strip() and (self.para is None or not self.para.runs):
                return
        if self.link is not None:
            self.link_text.append(data)
            return
        run = self._ensure().add_run(_clean(data))
        run.bold = "b" in self.fmt or None
        run.italic = "i" in self.fmt or None
        run.underline = "u" in self.fmt or None
        if "code" in self.fmt:
            _run_font(run, MONO)
            run.font.size = Pt(9)
        if "sub" in self.fmt:
            run.font.subscript = True
        if "sup" in self.fmt:
            run.font.superscript = True


def _burp_html(container, fragment: str, body_font: str) -> None:
    conv = _HtmlToDocx(container, body_font)
    conv.feed(sanitize_html(fragment or ""))
    conv.close()


# ---------------------------------------------------------------- the document

class _FastDoc:
    """A python-docx Document whose add_paragraph/add_heading/add_table insert in constant time.

    python-docx looks up the section properties by scanning the whole body on every insert, which makes
    large reports quadratic (minutes for a big scan). Everything else is forwarded to the real document.
    """

    def __init__(self, doc):
        self._doc = doc
        self._sect = doc.element.body.sectPr

    def __getattr__(self, name):
        return getattr(self._doc, name)

    def insert(self, el):
        self._sect.addprevious(el)
        return el

    def add_paragraph(self, text: str = "", style=None):
        p = Paragraph(self.insert(OxmlElement("w:p")), self._doc._body)
        if text:
            p.add_run(text)
        if style is not None:
            p.style = style
        return p

    def add_heading(self, text: str = "", level: int = 1):
        return self.add_paragraph(text, "Title" if level == 0 else f"Heading {level}")

    def add_table(self, rows: int, cols: int):
        return Table(self.insert(CT_Tbl.new_tbl(rows, cols, Cm(16.6))), self._doc._body)


class _Builder:
    def __init__(self, ctx: dict, theme: dict):
        self.ctx = ctx
        self.t = theme
        self.doc = _FastDoc(Document())
        self.bid = 0
        self.sec_no = 0
        self._code_proto = None
        self._setup()

    # ---------------------------------------------- setup
    def _setup(self):
        doc, t = self.doc, self.t
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
        styles = doc.styles
        normal = styles["Normal"]
        _set_font(normal, t["body"])
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = _rgb(TEXT)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.12
        for name, size, before in (("Heading 1", 18, 18), ("Heading 2", 14, 14), ("Heading 3", 12, 10)):
            st = styles[name]
            _set_font(st, t["heading"])
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = _rgb(t["accent"])
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        title = styles["Title"]
        _set_font(title, t["heading"])
        title.font.size = Pt(30)
        title.font.color.rgb = _rgb(t["accent"])
        for name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
            if name in [s.name for s in styles]:
                _set_font(styles[name], t["body"])
        # Ask Word to refresh fields (table of contents, page numbers) when the file is opened.
        settings = doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
        self._footer(sec)

    def _footer(self, sec):
        meta = self.ctx["meta"]
        sec.different_first_page_header_footer = bool(self.ctx["sections"].get("cover"))
        p = sec.footer.paragraphs[0]
        width = sec.page_width - sec.left_margin - sec.right_margin
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        left = p.add_run(_clean(meta["classification"] or meta["footer"] or meta["title"]))
        left.font.size = Pt(8)
        left.font.color.rgb = _rgb(MUTED)
        r = p.add_run("\tPage ")
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(MUTED)
        _field(p, "PAGE", "1")
        r = p.add_run(" of ")
        r.font.size = Pt(8)
        r.font.color.rgb = _rgb(MUTED)
        _field(p, "NUMPAGES", "1")
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = _rgb(MUTED)

    # ---------------------------------------------- small building blocks
    def heading(self, text: str, level: int = 1, anchor: str | None = None, page_break: bool = False):
        if level == 1 and self.t["numbered"]:
            self.sec_no += 1
            text = f"{self.sec_no}. {text}"
        h = self.doc.add_heading(_clean(text), level=level)
        if page_break:
            h.paragraph_format.page_break_before = True
        if anchor:
            self.bid += 1
            _bookmark(h, anchor, self.bid)
        return h

    def para(self, text: str = "", bold=False, italic=False, size=None, color=None, after=None, style=None):
        p = self.doc.add_paragraph(style=style)
        if text:
            r = p.add_run(_clean(text))
            r.bold = bold or None
            r.italic = italic or None
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = _rgb(color)
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def paragraphs(self, text: str):
        for block in [b.strip() for b in re.split(r"\n\s*\n", text or "") if b.strip()]:
            p = self.doc.add_paragraph()
            for n, line in enumerate(block.split("\n")):
                if n:
                    p.add_run().add_break()
                p.add_run(_clean(line))

    def badge(self, paragraph, sev: str, size: float = 8.5):
        r = paragraph.add_run(f" {_clean(sev)} ")
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(SEV_TEXT.get(sev, "FFFFFF"))
        _run_shade(r, SEV_COLORS.get(sev, SEV_COLORS["Information"]))
        return r

    def subhead(self, text: str):
        p = self.para(text.upper(), bold=True, size=8.5, color=MUTED, after=2)
        p.paragraph_format.space_before = Pt(8)
        _keep_with_next(p)
        return p

    def table(self, headers: list[str], rows: list[list], widths: list[float] | None = None, header_fill=None):
        tbl = self.doc.add_table(rows=1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(tbl)
        _cell_margins(tbl)
        hdr = tbl.rows[0]
        _repeat_header(hdr)
        fill = header_fill or self.t["accent"]
        for cell, text in zip(hdr.cells, headers):
            _shade(cell, fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(_clean(text))
            r.bold = True
            r.font.size = Pt(9)
            r.font.color.rgb = _rgb("FFFFFF")
        for values in rows:
            row = tbl.add_row()
            _no_split(row)
            for cell, value in zip(row.cells, values):
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if callable(value):
                    value(p)
                else:
                    r = p.add_run(_clean(value))
                    r.font.size = Pt(9)
        if widths:
            _widths(tbl, widths)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return tbl

    def kv_table(self, rows: list[tuple[str, object]], label_width: float = 4.2, total: float = 16.6):
        tbl = self.doc.add_table(rows=0, cols=2)
        _borders(tbl)
        _cell_margins(tbl)
        for label, value in rows:
            row = tbl.add_row()
            _no_split(row)
            lc, vc = row.cells
            _shade(lc, SURFACE)
            lp = lc.paragraphs[0]
            lp.paragraph_format.space_after = Pt(0)
            lr = lp.add_run(_clean(label))
            lr.bold = True
            lr.font.size = Pt(9)
            lr.font.color.rgb = _rgb(MUTED)
            vp = vc.paragraphs[0]
            vp.paragraph_format.space_after = Pt(0)
            if callable(value):
                value(vp)
            else:
                vr = vp.add_run(_clean(value))
                vr.font.size = Pt(9.5)
        _widths(tbl, [label_width, total - label_width])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return tbl

    def code_block(self, title: str, text: str, marks: list[tuple[int, int]]):
        self.subhead(title)
        end = len(text.rstrip())
        text = text[:end]
        marks = [(s, min(e, end)) for s, e in marks if s < end]
        if self._code_proto is None:
            # Built once with python-docx, then cloned: creating tables through the API is slow on big scans.
            tbl = self.doc.add_table(rows=1, cols=1)
            _borders(tbl, BORDER)
            _cell_margins(tbl, 100, 100, 140, 140)
            _widths(tbl, [16.6])
            cell = tbl.rows[0].cells[0]
            _shade(cell, CODE_BG)
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.0
            self._code_proto = tbl._tbl
            tbl._tbl.getparent().remove(tbl._tbl)
        el = self.doc.insert(copy.deepcopy(self._code_proto))
        p = el.find(".//" + qn("w:p"))
        pos = 0
        segments = []
        for s, e in marks:
            segments.append((text[pos:s], False))
            segments.append((text[s:e], True))
            pos = e
        segments.append((text[pos:], False))
        # Built as raw XML: one run per segment with <w:br/> between lines. Going through python-docx's
        # run API per line is ~50x slower on large scans.
        for seg, hl in segments:
            if seg:
                p.append(_code_run(seg, hl))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def link_or_text(self, p, text: str, anchor: str, linked: bool, bold=True):
        if linked:
            _hyperlink(p, text, anchor=anchor, color=self.t["accent"], bold=bold)
        else:
            r = p.add_run(_clean(text))
            r.bold = bold or None
            r.font.size = Pt(9)

    # ---------------------------------------------- sections
    def build(self) -> bytes:
        c, s = self.ctx, self.ctx["sections"]
        self.cover()
        if s.get("toc"):
            self.toc()
        first_after_cover = [s.get("cover") and not s.get("toc")]

        def pb():
            # The first section after the cover starts on a new page (the TOC already ends with a break).
            if first_after_cover[0]:
                first_after_cover[0] = False
                return True
            return False

        if s.get("executive_summary"):
            self.heading("Executive summary", anchor="executive_summary", page_break=pb())
            self.severity_tiles()
            self.paragraphs(c["executive_summary"])
        if s.get("scope") and c["scope"]:
            self.heading("Scope & methodology", anchor="scope", page_break=pb())
            self.paragraphs(c["scope"])
        if s.get("risk_chart") or s.get("findings_table"):
            self.heading("Findings overview", anchor="overview", page_break=pb())
            if s.get("risk_chart"):
                self.risk_chart()
            if s.get("findings_table"):
                self.findings_table()
        if s.get("hosts_table"):
            self.heading("Affected hosts", anchor="hosts", page_break=pb())
            self.hosts_table()
        if s.get("owasp_table"):
            self.heading("OWASP Top 10 mapping", anchor="owasp", page_break=pb())
            self.owasp_table()
        if s.get("findings"):
            self.heading("Detailed findings", anchor="findings", page_break=True)
            pb()
            if not c["findings"]:
                self.para("No findings match the selected filters.", italic=True, color=MUTED)
            for n, f in enumerate(c["findings"]):
                self.finding(f, page_break=c["cfg"].page_break_per_finding and n > 0)
        if s.get("scan_info"):
            self.heading("Appendix: scan information", anchor="scan_info", page_break=pb())
            self.appendix()
        if c["meta"]["footer"]:
            self.para(c["meta"]["footer"], size=9, color=MUTED).paragraph_format.space_before = Pt(18)
        out = io.BytesIO()
        self.doc.save(out)
        return out.getvalue()

    def cover(self):
        meta, doc = self.ctx["meta"], self.doc
        if not self.ctx["sections"].get("cover"):
            self.para(meta["title"], bold=True, size=20, color=self.t["accent"], after=2)
            bits = [b for b in (meta["client"], meta["date"], meta["classification"]) if b]
            self.para(" · ".join(bits), color=MUTED, after=12)
            return
        self._logo()
        for _ in range(4):
            doc.add_paragraph()
        if meta["classification"]:
            r = self.para(after=10).add_run(_clean(meta["classification"]))
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = _rgb(SEV_COLORS["High"])
        doc.add_paragraph(_clean(meta["title"]), style="Title")
        if meta["subtitle"]:
            self.para(meta["subtitle"], size=15, color=MUTED, after=14)
        if meta["client"]:
            p = self.para(after=24)
            r = p.add_run("Prepared for ")
            r.font.size = Pt(13)
            r = p.add_run(_clean(meta["client"]))
            r.bold = True
            r.font.size = Pt(13)
        rows = [(k, v) for k, v in (("Prepared by", meta["prepared_by"]), ("Organisation", meta["company"]),
                                    ("Date", meta["date"]), ("Version", meta["version"])) if v]
        tbl = doc.add_table(rows=0, cols=2)
        _no_borders(tbl)
        for k, v in rows:
            row = tbl.add_row()
            a, b = row.cells
            a.width, b.width = Cm(3.5), Cm(9)
            ra = a.paragraphs[0].add_run(k)
            ra.font.color.rgb = _rgb(MUTED)
            rb = b.paragraphs[0].add_run(_clean(v))
            rb.bold = True
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _logo(self):
        logo = self.ctx["meta"]["logo"]
        m = re.match(r"data:image/(png|jpeg|gif);base64,(.+)$", logo or "")
        if not m:
            return  # SVG/WebP logos aren't supported by Word's picture import
        try:
            self.doc.add_picture(io.BytesIO(base64.b64decode(m.group(2))), height=Cm(1.8))
        except Exception:  # noqa: BLE001 - a broken image shouldn't stop the export
            pass

    def toc(self):
        p = self.para("Contents", bold=True, size=18, color=self.t["accent"], after=8)
        _keep_with_next(p)
        entries = []
        s, c = self.ctx["sections"], self.ctx
        n = 0
        for key, title, show in (
            ("executive_summary", "Executive summary", s.get("executive_summary")),
            ("scope", "Scope & methodology", s.get("scope") and c["scope"]),
            ("overview", "Findings overview", s.get("risk_chart") or s.get("findings_table")),
            ("hosts", "Affected hosts", s.get("hosts_table")),
            ("owasp", "OWASP Top 10 mapping", s.get("owasp_table")),
            ("findings", "Detailed findings", s.get("findings")),
            ("scan_info", "Appendix: scan information", s.get("scan_info")),
        ):
            if show:
                n += 1
                entries.append((f"{n}. {title}" if self.t["numbered"] else title, key, 0))
                if key == "findings":
                    entries += [(f"{f.id}  {f.name}", f.anchor.replace("-", "_"), 1) for f in c["findings"]]
        # A real TOC field (Word fills in page numbers when it updates fields on open); until then it
        # shows this static list with working links.
        tp = self.doc.add_paragraph()
        _field_toc(tp, entries, self.t["accent"])
        note = self.para("Page numbers appear when Word updates fields (it asks when the file is opened; "
                         "or right-click the list and choose Update Field).", size=8, italic=True, color=MUTED)
        note.add_run().add_break(WD_BREAK.PAGE)

    def severity_tiles(self):
        sevs, counts = self.ctx["severities"], self.ctx["counts"]
        if not sevs:
            return
        tbl = self.doc.add_table(rows=1, cols=len(sevs))
        _borders(tbl, "FFFFFF", size=12)
        _cell_margins(tbl, 120, 120, 140, 140)
        for cell, sev in zip(tbl.rows[0].cells, sevs):
            _shade(cell, SEV_COLORS.get(sev, SEV_COLORS["Information"]))
            ink = _rgb(SEV_TEXT.get(sev, "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(counts[sev]))
            r.bold = True
            r.font.size = Pt(22)
            r.font.color.rgb = ink
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(0)
            r2 = p2.add_run(sev)
            r2.font.size = Pt(9)
            r2.font.color.rgb = ink
        self.doc.add_paragraph()

    def risk_chart(self):
        c = self.ctx
        sevs, counts, total = c["severities"], c["counts"], max(c["total_findings"], 1)
        biggest = max([counts[s] for s in sevs] + [1])

        def bar(sev):
            def draw(p):
                n = counts[sev]
                if n:
                    r = p.add_run("█" * max(1, round(n / biggest * 28)))
                    r.font.color.rgb = _rgb(SEV_COLORS.get(sev, SEV_COLORS["Information"]))
                    r.font.size = Pt(9)
                r = p.add_run(f"  {round(n / total * 100)}%")
                r.font.size = Pt(8.5)
                r.font.color.rgb = _rgb(MUTED)
            return draw

        self.table(["Severity", "Findings", "Instances", "Share of findings"],
                   [[lambda p, s=s: self.badge(p, s), str(counts[s]), str(c["instance_counts"][s]), bar(s)] for s in sevs],
                   widths=[3.2, 2.2, 2.4, 8.8])
        self.para(f"{c['total_findings']} unique finding{'s' if c['total_findings'] != 1 else ''} "
                  f"({c['total_instances']} instance{'s' if c['total_instances'] != 1 else ''}) after filtering.",
                  size=9, color=MUTED)

    def findings_table(self):
        c = self.ctx
        cwe = c["fields"].get("cwe_owasp")
        linked = c["sections"].get("findings")
        headers = ["ID", "Finding"] + (["CWE"] if cwe else []) + ["Severity", "Confidence", "Inst."]
        rows = []
        for f in c["findings"]:
            row = [lambda p, f=f: self.link_or_text(p, f.id, f.anchor.replace("-", "_"), linked), f.name]
            if cwe:
                row.append(", ".join(f"CWE-{i}" for i, _ in f.cwes[:2]) or "—")
            row += [lambda p, f=f: self.badge(p, f.severity), f.confidence, str(f.count)]
            rows.append(row)
        if not rows:
            self.para("No findings match the selected filters.", italic=True, color=MUTED)
            return
        widths = [1.7, 6.4, 2.4, 2.4, 2.4, 1.3] if cwe else [1.8, 8.7, 2.4, 2.4, 1.3]
        self.table(headers, rows, widths)

    def hosts_table(self):
        c = self.ctx
        sevs = c["severities"]
        rows = []
        for h in c["hosts"]:
            rows.append([h["host"] or "—", h["ip"] or "—"] + [str(h["counts"][s]) if h["counts"][s] else "–" for s in sevs]
                        + [str(h["total"])])
        if not rows:
            self.para("No hosts.", italic=True, color=MUTED)
            return
        sev_w = [min(2.2, 6.6 / max(len(sevs), 1))] * len(sevs)
        self.table(["Host", "IP address"] + [SEV_SHORT.get(x, x) for x in sevs] + ["Total"], rows, _spread([], 1, 16.6 - 2.8 - sum(sev_w) - 1.4)
                   + [2.8] + sev_w + [1.4])
        self.para("Counts are per issue instance.", size=9, color=MUTED)

    def owasp_table(self):
        c = self.ctx
        o, sevs, linked = c["owasp"], c["severities"], c["sections"].get("findings")
        self.para(f"Findings mapped to the {o['edition']} categories using the CWE identifiers in Burp's export. "
                  f"A finding can map to more than one category. {o['mapped_total']} of {c['total_findings']} "
                  f"finding{'s' if c['total_findings'] != 1 else ''} mapped.", size=9, color=MUTED)

        def ids(findings):
            def draw(p):
                if not findings:
                    r = p.add_run("None")
                    r.font.size = Pt(9)
                    r.font.color.rgb = _rgb(MUTED)
                for n, f in enumerate(findings):
                    if n:
                        p.add_run(", ").font.size = Pt(9)
                    self.link_or_text(p, f.id, f.anchor.replace("-", "_"), linked)
            return draw

        rows = []
        for r in o["rows"]:
            rows.append([f"{r['code']} {r['name']}"] + [str(r["counts"][s]) if r["counts"][s] else "–" for s in sevs]
                        + [ids(r["findings"])])
        if o["unmapped"]:
            rows.append(["Not mapped"] + [str(o["unmapped_counts"][s]) if o["unmapped_counts"][s] else "–" for s in sevs]
                        + [ids(o["unmapped"])])
        sev_w = [min(2.1, 8.4 / max(len(sevs), 1))] * len(sevs)
        self.table(["Category"] + [SEV_SHORT.get(x, x) for x in sevs] + ["Findings"], rows, [5.6] + sev_w + [max(16.6 - 5.6 - sum(sev_w), 2.5)])

    def finding(self, f, page_break: bool):
        c, fields = self.ctx, self.ctx["fields"]
        h = self.doc.add_heading(level=2)
        if page_break:
            h.paragraph_format.page_break_before = True
        h.add_run(_clean(f"{f.id}  {f.name}"))
        self.bid += 1
        _bookmark(h, f.anchor.replace("-", "_"), self.bid)
        meta = [("Severity", lambda p: self.badge(p, f.severity, 9)), ("Confidence", f.confidence),
                ("Instances", str(f.count)), ("Host" + ("s" if len(f.hosts) > 1 else ""),
                                             ", ".join(h for h in f.hosts if h) or "—")]
        if fields.get("cwe_owasp") and f.cwes:
            meta.append(("CWE", "\n".join(f"CWE-{i}" + (f": {t}" if t else "") for i, t in f.cwes)))
        if fields.get("cwe_owasp") and f.owasp:
            meta.append(("OWASP Top 10", "\n".join(f"{code}:2021 {name}" for code, name in f.owasp)))
        if f.type_id:
            meta.append(("Burp issue type", f.type_id))
        self.kv_table([(k, v if callable(v) else _multiline(v)) for k, v in meta])

        if fields.get("issue_background") and f.issue_background:
            self.heading("Description", level=3)
            _burp_html(self.doc, f.issue_background, self.t["body"])
        if fields.get("affected_urls"):
            self.heading("Affected URLs", level=3)
            rows = [[str(n), lambda p, u=i.url: _mono(p, u), lambda p, s=i.severity: self.badge(p, s), i.confidence]
                    for n, i in enumerate(f.instances, 1)]
            if f.omitted_instances:
                rows.append(["", f"… and {f.omitted_instances} more instance{'s' if f.omitted_instances > 1 else ''} not shown", "", ""])
            self.table(["#", "URL", "Severity", "Confidence"], rows, [0.9, 11.0, 2.4, 2.3], header_fill=MUTED)
        for n, inst in enumerate(f.instances, 1):
            body = ((fields.get("issue_detail") and inst.issue_detail) or (fields.get("detail_items") and inst.detail_items)
                    or (fields.get("remediation_detail") and inst.remediation_detail) or inst.evidence)
            if not body:
                continue
            if len(f.instances) > 1:
                p = self.para(after=4)
                p.paragraph_format.space_before = Pt(10)
                r = p.add_run(f"Instance {n}  ")
                r.bold = True
                r.font.color.rgb = _rgb(self.t["accent"])
                _mono(p, inst.url)
                _keep_with_next(p)
            if fields.get("issue_detail") and inst.issue_detail:
                self.subhead("Issue detail")
                _burp_html(self.doc, inst.issue_detail, self.t["body"])
            if fields.get("detail_items") and inst.detail_items:
                self.subhead("Detail items")
                for d in inst.detail_items:
                    self.doc.add_paragraph(_clean(d), style="List Bullet")
            if fields.get("remediation_detail") and inst.remediation_detail:
                self.subhead("Remediation detail")
                _burp_html(self.doc, inst.remediation_detail, self.t["body"])
            many = len(inst.evidence) > 1
            for k, ev in enumerate(inst.evidence, 1):
                num = f" {k}" if many else ""
                if ev["request_text"]:
                    note = " (truncated)" if ev["request_truncated"] else ""
                    note += " · payload highlighted" if ev["request_marks"] else ""
                    self.code_block(f"Request{num}{note}", ev["request_text"], ev["request_marks"])
                if ev["response_text"]:
                    note = " (after redirect)" if ev["redirected"] else ""
                    note += " (headers only)" if ev["response_headers_only"] else (" (truncated)" if ev["response_truncated"] else "")
                    note += " · payload highlighted" if ev["response_marks"] else ""
                    self.code_block(f"Response{num}{note}", ev["response_text"], ev["response_marks"])
        if fields.get("remediation_background") and f.remediation_background:
            self.heading("Remediation", level=3)
            _burp_html(self.doc, f.remediation_background, self.t["body"])
        if fields.get("references") and f.references:
            self.heading("References", level=3)
            _burp_html(self.doc, f.references, self.t["body"])
        if fields.get("classifications") and f.classifications:
            self.heading("Vulnerability classifications", level=3)
            _burp_html(self.doc, f.classifications, self.t["body"])

    def appendix(self):
        c = self.ctx
        sc, red = c["scan"], c["redaction"]
        rows = [("Source file", sc["source"])]
        if sc["burp_version"]:
            rows.append(("Burp Suite version", sc["burp_version"]))
        if sc["export_time"]:
            rows.append(("Export time", sc["export_time"]))
        rows += [("Issues in export", str(sc["total_issues"])),
                 ("Issues included", f"{c['total_instances']}"
                  + (f" ({sc['filtered_out']} excluded by filters)" if sc["filtered_out"] else "")),
                 ("Severities included", ", ".join(c["severities"])),
                 ("Confidence levels included", ", ".join(c["cfg"].confidences))]
        if c["fields"].get("evidence"):
            rows.append(("Evidence redaction",
                         f"Enabled: {red['count']} sensitive value{'s' if red['count'] != 1 else ''} masked as [REDACTED]"
                         if red["enabled"] or red["custom"] else "Disabled"))
        rows.append(("Report generated", c["generated_at"]))
        self.kv_table(rows, label_width=5.2)
        self.heading("Severity ratings", level=3)
        self.kv_table([
            ("High", "Issues that could lead to significant compromise of data, systems or users and should be fixed as a priority."),
            ("Medium", "Issues that weaken security and may be exploitable in combination with other weaknesses."),
            ("Low", "Issues with limited direct impact that represent deviations from best practice."),
            ("Information", "Observations that do not represent a vulnerability on their own but may be of interest."),
        ], label_width=3.2)
        self.heading("Confidence levels", level=3)
        self.kv_table([
            ("Certain", "Burp is certain the issue exists."),
            ("Firm", "The issue is very likely present, but has not been fully confirmed."),
            ("Tentative", "The issue may be present; manual verification is recommended."),
        ], label_width=3.2)


def _code_run(text: str, highlighted: bool):
    """A monospace run; newlines become <w:br/>. Highlighted runs are bold on yellow."""
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), MONO)
    rpr.append(fonts)
    if highlighted:
        rpr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")  # half-points: 8 pt
    rpr.append(sz)
    if highlighted:
        hl = OxmlElement("w:highlight")
        hl.set(qn("w:val"), "yellow")
        rpr.append(hl)
    run.append(rpr)
    # etree.SubElement rather than OxmlElement, which parses an XML string per element (slow for big evidence).
    tag_t, tag_br, space = qn("w:t"), qn("w:br"), qn("xml:space")
    for n, line in enumerate(_clean(text).split("\n")):
        if n:
            etree.SubElement(run, tag_br)
        if line:
            t = etree.SubElement(run, tag_t)
            t.set(space, "preserve")
            t.text = line
    return run


def _mono(p, text: str):
    r = p.add_run(_clean(text))
    _run_font(r, MONO)
    r.font.size = Pt(8.5)
    return r


def _multiline(text: str):
    def draw(p):
        for n, line in enumerate(str(text).split("\n")):
            if n:
                p.add_run().add_break()
            r = p.add_run(_clean(line))
            r.font.size = Pt(9.5)
    return draw


def _field_toc(paragraph, entries: list[tuple[str, str, int]], accent: str) -> None:
    """A TOC field whose cached result is a linked list of the report's headings."""
    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            el.set(qn("w:dirty"), "true")
        return el
    r = paragraph.add_run()
    r._element.append(fld("begin"))
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._element.append(it)
    r = paragraph.add_run()
    r._element.append(fld("separate"))
    for n, (text, anchor, level) in enumerate(entries):
        if n:
            paragraph.add_run().add_break()
        if level:
            paragraph.add_run("      ")
        _hyperlink(paragraph, text, anchor=anchor, color=accent if not level else TEXT, bold=not level)
    r = paragraph.add_run()
    r._element.append(fld("end"))


def render_docx(scan: ScanReport, cfg: ReportConfig) -> bytes:
    if not DOCX_AVAILABLE:
        raise RuntimeError(DOCX_MISSING)
    ctx = build_context(scan, cfg, "docx")
    return _Builder(ctx, _theme(cfg)).build()
